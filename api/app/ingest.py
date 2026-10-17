"""Turn uploaded business files into clean tables (CSV/TSV/Excel/JSON) or searchable text (PDF/DOCX/TXT/MD)."""
import csv
import io
import json
import re
import warnings
from pathlib import Path

import pandas as pd

from .config import get_settings

TABULAR_EXT = {".csv", ".tsv", ".xlsx", ".xlsm", ".xls", ".json", ".jsonl"}
DOC_EXT = {".pdf", ".docx", ".txt", ".md"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
RESERVED = {"order", "group", "user", "table", "select", "from", "where", "index", "values", "default",
            "limit", "join", "union", "all", "end", "case", "check", "column", "key", "primary", "desc", "asc"}

_ID_LIKE = re.compile(r"(^|_)(id|code|zip|pin|pincode|phone|mobile|sku|gstin|pan|ifsc|account|invoice|number|no|"
                      r"barcode|isbn|hsn)($|_)")
_DATE_NAME = re.compile(r"(date|time|dt|day|month|period|created|updated|timestamp)")
_DATE_VAL = re.compile(r"^\s*(\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})([ T].*)?$")


def kind_of(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in TABULAR_EXT:
        return "table"
    if ext in DOC_EXT:
        return "document"
    if ext in IMAGE_EXT:
        return "image"
    return "unsupported"


# ------------------------------------------------------------------- naming
def snake(name, i: int) -> str:
    s = "" if name is None or (isinstance(name, float) and pd.isna(name)) else str(name).strip()
    s = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", s)
    s = re.sub(r"[^\w]+", "_", s).strip("_").lower()
    if not s or s.startswith("unnamed"):
        s = f"col_{i + 1}"
    if s[0].isdigit():
        s = "c_" + s
    return s[:60]


def unique(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        if n in seen:
            seen[n] += 1
            out.append(f"{n}_{seen[n]}")
        else:
            seen[n] = 1
            out.append(n)
    return out


def table_name(filename: str, sheet: str | None = None, taken: set[str] | None = None) -> str:
    stem = Path(filename).stem
    base = snake(stem + ("_" + sheet if sheet else ""), 0)
    base = re.sub(r"_+", "_", base)[:40].strip("_") or "data"
    if base[0].isdigit():
        base = "t_" + base
    if base in RESERVED:
        base += "_data"
    taken = taken or set()
    name, n = base, 1
    while name in taken:
        n += 1
        name = f"{base}_{n}"
    return name


# ------------------------------------------------------------------- reading
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_csv(data: bytes, filename: str) -> pd.DataFrame:
    text = _decode(data)
    sample = text[:20000]
    delim = "\t" if filename.lower().endswith(".tsv") else ","
    if not filename.lower().endswith(".tsv"):
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            pass
    return pd.read_csv(io.StringIO(text), sep=delim, low_memory=False, on_bad_lines="skip", dtype=str,
                       nrows=get_settings().max_table_rows + 1)


def _promote_header(raw: pd.DataFrame) -> pd.DataFrame:
    raw = raw.dropna(how="all").dropna(axis=1, how="all")
    if raw.empty:
        return raw
    raw = raw.reset_index(drop=True)
    maxn = int(raw.notna().sum(axis=1).max())
    pick = 0
    for i in range(min(15, len(raw))):
        row = raw.iloc[i]
        nn = int(row.notna().sum())
        strs = sum(isinstance(v, str) for v in row if pd.notna(v))
        if nn >= 0.6 * maxn and strs >= 0.6 * nn:
            pick = i
            break
    header = list(raw.iloc[pick])
    body = raw.iloc[pick + 1:].reset_index(drop=True)
    body.columns = header
    return body


def _read_excel(data: bytes) -> list[tuple[str, pd.DataFrame]]:
    xl = pd.ExcelFile(io.BytesIO(data))
    out = []
    for sheet in xl.sheet_names:
        raw = xl.parse(sheet, header=None, nrows=get_settings().max_table_rows + 20)
        df = _promote_header(raw)
        if not df.empty and len(df.columns):
            out.append((sheet, df))
    return out


def _read_json(data: bytes, filename: str) -> pd.DataFrame:
    text = _decode(data)
    if filename.lower().endswith(".jsonl"):
        return pd.read_json(io.StringIO(text), lines=True)
    obj = json.loads(text)
    if isinstance(obj, dict):
        lists = [v for v in obj.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        obj = lists[0] if lists else [obj]
    return pd.json_normalize(obj)


# ------------------------------------------------------------------- cleaning
def _dayfirst(values: pd.Series) -> bool:
    a = b = 0
    for v in values.head(200):
        m = re.match(r"\s*(\d{1,2})[-/.](\d{1,2})[-/.]\d{2,4}", str(v))
        if m:
            a += int(m.group(1)) > 12
            b += int(m.group(2)) > 12
    return a > b


def clean_frame(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """Normalise names, drop empty rows/cols, fix numbers stored as text, parse dates."""
    s = get_settings()
    notes: list[str] = []
    df = df.dropna(how="all").dropna(axis=1, how="all").copy()
    if len(df.columns) > s.max_columns:
        notes.append(f"Only the first {s.max_columns} columns were kept.")
        df = df.iloc[:, : s.max_columns]
    if len(df) > s.max_table_rows:
        notes.append(f"Only the first {s.max_table_rows:,} rows were kept.")
        df = df.iloc[: s.max_table_rows]
    df.columns = unique([snake(c, i) for i, c in enumerate(df.columns)])

    for col in df.columns:
        ser = df[col]
        if not (ser.dtype == object or pd.api.types.is_string_dtype(ser)):
            continue
        nn = ser.dropna()
        if nn.empty:
            continue
        st = nn.astype(str).str.strip()
        idlike = bool(_ID_LIKE.search(col)) or st.str.match(r"^0\d+$").any()
        if not idlike:
            cleaned = (st.str.replace(r"[,\s₹$€£]", "", regex=True)
                         .str.replace(r"^\((.*)\)$", r"-\1", regex=True)
                         .str.replace(r"%$", "", regex=True))
            num = pd.to_numeric(cleaned, errors="coerce")
            if num.notna().mean() >= 0.97 and not (cleaned.str.len() > 15).any():
                full = pd.to_numeric(ser.astype(str).str.strip()
                                     .str.replace(r"[,\s₹$€£]", "", regex=True)
                                     .str.replace(r"^\((.*)\)$", r"-\1", regex=True)
                                     .str.replace(r"%$", "", regex=True), errors="coerce")
                df[col] = full
                continue
        looks_date = bool(_DATE_NAME.search(col)) or st.head(50).str.match(_DATE_VAL).mean() >= 0.8
        if looks_date and not idlike:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                parsed = pd.to_datetime(ser, errors="coerce", dayfirst=_dayfirst(st), format="mixed")
            if parsed.notna().sum() >= 0.9 * len(nn):
                df[col] = parsed
                continue
        # keep as text; make sure it is uniformly str for DuckDB
        df[col] = ser.map(lambda v: None if pd.isna(v) else str(v))
    return df.reset_index(drop=True), notes


def parse_tabular(filename: str, data: bytes) -> list[tuple[str | None, pd.DataFrame, list[str]]]:
    """Return [(sheet_name|None, cleaned_frame, notes)]."""
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".tsv"):
        frames = [(None, _read_csv(data, filename))]
    elif ext in (".xlsx", ".xlsm", ".xls"):
        frames = _read_excel(data)
    elif ext in (".json", ".jsonl"):
        frames = [(None, _read_json(data, filename))]
    else:
        raise ValueError(f"Unsupported table format: {ext}")
    out = []
    for sheet, df in frames:
        clean, notes = clean_frame(df)
        if len(clean) and len(clean.columns):
            out.append((sheet, clean, notes))
    if not out:
        raise ValueError("The file has no readable rows. Check that it has a header row and data.")
    return out


# ------------------------------------------------------------------- documents
def parse_document(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md"):
        text = _decode(data)
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages[:300])
    elif ext == ".docx":
        from docx import Document

        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(f"Unsupported document format: {ext}")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError("No text could be extracted (scanned PDFs need OCR, which is not supported yet).")
    return text[:3_000_000]
